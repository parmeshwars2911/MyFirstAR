"""
Lightweight PPTX -> PNG previewer for visual QA.

Reads a saved .pptx and re-renders each slide to PNG via cairosvg by walking
the shape tree (auto-shapes with solid fills + text, and pictures). It is an
approximation good enough to spot overlaps, clipping, colour and layout
problems without needing LibreOffice.
"""
import os
import sys
import base64
import html
from pptx import Presentation
from pptx.util import Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
import cairosvg

EMU_IN = 914400
SCALE = 96  # px per inch


def _bg(slide):
    bg = slide._element.find(qn("p:cSld") + "/" + qn("p:bg"))
    if bg is not None:
        clr = bg.find(".//" + qn("a:srgbClr"))
        if clr is not None:
            return "#" + clr.get("val")
    return "#FFFFFF"


def _fill(sp):
    try:
        f = sp.fill
        if f.type == 1:
            return "#" + str(f.fore_color.rgb)
    except Exception:
        pass
    return None


def _line(sp):
    try:
        ln = sp.line
        if ln.color and ln.color.type is not None:
            return "#" + str(ln.color.rgb), max(1, (ln.width or 12700) / 12700)
    except Exception:
        pass
    return None, 0


def _runs(tf):
    out = []
    for p in tf.paragraphs:
        align = {PP_ALIGN.CENTER: "middle", PP_ALIGN.RIGHT: "end"}.get(
            p.alignment, "start")
        runs = []
        for r in p.runs:
            f = r.font
            runs.append(dict(
                t=r.text,
                size=(f.size.pt if f.size else 15),
                bold=bool(f.bold),
                italic=bool(f.italic),
                color="#" + str(f.color.rgb) if (f.color and f.color.type
                                                 is not None) else "#1D2433",
            ))
        out.append((align, runs))
    return out


def _anchor(tf):
    try:
        return {MSO_ANCHOR.MIDDLE: "middle", MSO_ANCHOR.BOTTOM: "bottom"}.get(
            tf.vertical_anchor, "top")
    except Exception:
        return "top"


def render_pptx(path, outdir, slides=None, w_px=1066):
    prs = Presentation(path)
    sw = prs.slide_width / EMU_IN
    sh = prs.slide_height / EMU_IN
    h_px = int(w_px * sh / sw)
    sc = w_px / sw
    os.makedirs(outdir, exist_ok=True)
    base = os.path.splitext(os.path.basename(path))[0]
    paths = []
    for i, slide in enumerate(prs.slides, 1):
        if slides and i not in slides:
            continue
        body = [f'<rect width="{w_px}" height="{h_px}" fill="{_bg(slide)}"/>']
        for shp in slide.shapes:
            L = (shp.left or 0) / EMU_IN * sc
            T = (shp.top or 0) / EMU_IN * sc
            W = (shp.width or 0) / EMU_IN * sc
            H = (shp.height or 0) / EMU_IN * sc
            # picture
            if shp.shape_type == 13:
                try:
                    blob = shp.image.blob
                    ext = shp.image.ext
                    b64 = base64.b64encode(blob).decode()
                    body.append(
                        f'<image x="{L}" y="{T}" width="{W}" height="{H}" '
                        f'href="data:image/{ext};base64,{b64}"/>')
                except Exception:
                    pass
                continue
            fill = _fill(shp)
            lc, lw = _line(shp)
            try:
                ast = str(shp.auto_shape_type)
            except Exception:
                ast = ""
            is_oval = "OVAL" in ast
            is_round = "ROUNDED" in ast
            if fill or lc:
                attrs = f'fill="{fill}"' if fill else 'fill="none"'
                if lc:
                    attrs += f' stroke="{lc}" stroke-width="{lw}"'
                if is_oval:
                    body.append(f'<ellipse cx="{L+W/2}" cy="{T+H/2}" '
                                f'rx="{W/2}" ry="{H/2}" {attrs}/>')
                else:
                    rx = 10 if is_round else 0
                    body.append(f'<rect x="{L}" y="{T}" width="{W}" '
                                f'height="{H}" rx="{rx}" {attrs}/>')
            # text (with greedy word-wrap within the box width)
            if shp.has_text_frame and shp.text_frame.text.strip():
                anchor = _anchor(shp.text_frame)
                paras = _runs(shp.text_frame)
                maxw = max(W - 6, 20)
                # build wrapped lines: list of (lineruns, lineheight)
                lines = []
                for align, runs in paras:
                    if not runs:
                        lines.append((align, [], 15 * sc / 72 * 1.25))
                        continue
                    cur = []
                    curw = 0
                    maxsz = max(r["size"] for r in runs)
                    for r in runs:
                        words = r["t"].split(" ")
                        for wi, word in enumerate(words):
                            token = word if wi == 0 else " " + word
                            gw = len(token) * r["size"] * sc / 72 * 0.5
                            if curw + gw > maxw and cur:
                                lines.append((align, cur, maxsz * sc / 72 * 1.22))
                                cur = []
                                curw = 0
                                token = word
                                gw = len(token) * r["size"] * sc / 72 * 0.5
                            cur.append(dict(r, t=token))
                            curw += gw
                    if cur:
                        lines.append((align, cur, maxsz * sc / 72 * 1.22))
                total = sum(lh for _, _, lh in lines)
                if anchor == "middle":
                    y = T + (H - total) / 2
                elif anchor == "bottom":
                    y = T + (H - total)
                else:
                    y = T

                def _span(r):
                    fs = r["size"] * sc / 72
                    fw = "700" if r["bold"] else "400"
                    it = ' font-style="italic"' if r["italic"] else ""
                    return (f'<tspan fill="{r["color"]}" '
                            f'font-size="{fs:.1f}" font-weight="{fw}"{it}>'
                            f'{html.escape(r["t"])}</tspan>')
                for align, runs, lh in lines:
                    y += lh * 0.8
                    if runs:
                        tx = L + W / 2 if align == "middle" else (
                            L + W if align == "end" else L)
                        spans = "".join(_span(r) for r in runs)
                        body.append(
                            f'<text x="{tx}" y="{y:.1f}" text-anchor="{align}" '
                            f'font-family="DejaVu Sans, Arial">{spans}</text>')
                    y += lh * 0.2
                # overflow marker (visual QA aid)
                if y > T + H + 4:
                    body.append(f'<rect x="{L}" y="{T}" width="{W}" '
                                f'height="{H}" fill="none" stroke="#E63946" '
                                f'stroke-width="2" stroke-dasharray="6 4"/>')
        svg = (f'<svg xmlns="http://www.w3.org/2000/svg" width="{w_px}" '
               f'height="{h_px}" viewBox="0 0 {w_px} {h_px}">'
               + "".join(body) + "</svg>")
        outp = os.path.join(outdir, f"{base}_s{i:02d}.png")
        cairosvg.svg2png(bytestring=svg.encode("utf-8"), write_to=outp)
        paths.append(outp)
    return paths


if __name__ == "__main__":
    src = sys.argv[1]
    outdir = sys.argv[2] if len(sys.argv) > 2 else "build/preview"
    sl = None
    if len(sys.argv) > 3:
        sl = [int(x) for x in sys.argv[3].split(",")]
    for p in render_pptx(src, outdir, slides=sl):
        print(p)
